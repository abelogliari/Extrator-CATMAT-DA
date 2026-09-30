"""
Extrator Compras.gov  —  só a extração que conecta no Compras.gov, sem interface.

Isolado do ExtratorCatmat.py: é o mesmo motor que baixa os Registros de Preço do
Portal de Compras Governamentais (dadosabertos.compras.gov.br), sem a interface
CustomTkinter, sem a aba BPS e sem a consolidação DW + DA. Precisa apenas de
requests, pandas e openpyxl.

Como no ExtratorCatmat (opção "Usar a descrição do BPS", marcada por padrão), o
descricaoItem de cada registro é trocado pela descrição do mesmo CATMAT no BPS
(apidadosabertos.saude.gov.br). --sem-descricao-bps desliga a troca.

    Classe -> PDMs -> CATMATs -> Registros de Preço

    python extrator_comprasgov.py codigos.xlsx                    -> CATMATs do arquivo
    python extrator_comprasgov.py codigos.xlsx -t PDM             -> PDMs do arquivo
    python extrator_comprasgov.py -c "451234;451235"              -> códigos avulsos
    python extrator_comprasgov.py --classes "6505;6515" -s saida  -> classes inteiras
    python extrator_comprasgov.py -h                              -> todas as opções

Grava na pasta de saída (-s, padrão: a atual) o dados_completos_extraidos_part1
(.xlsx ou .csv) — ou um classe_XXXX_part1 por classe — e o relatório de
integridade. Ctrl+C cancela e salva o que já foi baixado; um segundo Ctrl+C
encerra na hora.
"""

import re
import csv
import math
import os
import sys
import time
import signal
import argparse
import threading
from io import StringIO
from typing import List, Optional
from concurrent.futures import ThreadPoolExecutor, wait, FIRST_COMPLETED

import requests
import pandas as pd
from openpyxl import Workbook

# ══════════════════════════════════════════════════════════════════════════════
#  CONFIGURAÇÃO  —  ajuste aqui
# ══════════════════════════════════════════════════════════════════════════════
URL_BASE       = "https://dadosabertos.compras.gov.br"
TIMEOUT        = 120      # segundos por requisição
TAMANHO_PAGINA = 500      # máximo aceito pela Pesquisa de Preço
_MAX_PAGINAS   = 20000    # trava contra laço infinito se a API paginar sem fim
COTA_COMPRAS_POR_MINUTO = 90   # margem abaixo das ~100 medidas: um 429 custa até 45 s
WORKERS_COMPRAS         = 3    # 3 em paralelo cobrem a latência (~0,8 s) e enchem a cota
# ══════════════════════════════════════════════════════════════════════════════

requests.packages.urllib3.disable_warnings(
    requests.packages.urllib3.exceptions.InsecureRequestWarning
)

# Session compartilhada — reutiliza conexões TCP/TLS entre todas as requisições
# Evita o overhead de handshake (~200-400ms) a cada chamada
_http = requests.Session()
_http.verify = False
_http.headers.update({"Accept-Encoding": "gzip, deflate", "Connection": "keep-alive"})
_adapter = requests.adapters.HTTPAdapter(
    pool_connections=4, pool_maxsize=12, max_retries=0
)
_http.mount("https://", _adapter)
_http.mount("http://",  _adapter)

# =============================================================================
# CANCELAMENTO (Ctrl+C)  E  LOG
# -----------------------------------------------------------------------------
# O tratador de sinal só levanta a bandeira: nada de locks lá dentro, porque ele
# roda na thread principal no meio do que ela estiver fazendo. As threads leem
# a bandeira entre uma página e outra; as esperas são em fatias de 0,5 s porque
# no Windows o sinal só é tratado quando a thread principal volta a rodar Python.
# =============================================================================
_cancelado = False
_lock_log  = threading.Lock()


def _ao_ctrl_c(signum, frame):
    global _cancelado
    _cancelado = True
    signal.signal(signal.SIGINT, signal.default_int_handler)   # o 2º encerra na hora


def _dormir(segundos):
    """time.sleep que acorda no Ctrl+C."""
    fim = time.monotonic() + segundos
    while not _cancelado:
        resta = fim - time.monotonic()
        if resta <= 0:
            break
        time.sleep(min(resta, 0.5))


def _log(msg=""):
    """print com hora, seguro entre threads."""
    with _lock_log:
        print(f"{time.strftime('%H:%M:%S')}  {msg}" if msg else "", flush=True)


def _milhar(n: int) -> str:
    return f"{n:,}".replace(",", ".")

# =============================================================================
# COTA DO COMPRAS.GOV
# -----------------------------------------------------------------------------
# O Compras.gov fica atrás de um Azure API Management que limita cada cliente a
# ~100 requisições por MINUTO (medido em 29/09/2026: 97-100 passam, depois vem
# 429 "Rate limit is exceeded. Try again in N seconds" com Retry-After de até
# 45 s). A cota é UMA só para todos os módulos: pesquisa de preço e catálogo
# (PDMs, itens) gastam do mesmo saldo.
#
# Por isso a velocidade máxima é a cota, não o número de threads: todas as
# chamadas ao Compras.gov passam por _get_compras, que espaça as saídas para
# ficar logo abaixo do limite e, se ainda assim vier um 429, segura TODAS as
# threads pelo tempo que o servidor pediu (Retry-After) e repete.
# =============================================================================


class _CotaCompras:
    """Espaçamento global das requisições + pausa global em caso de 429."""

    def __init__(self, por_minuto):
        self._intervalo = 60.0 / por_minuto
        self._lock      = threading.Lock()
        self._proximo   = 0.0          # instante mínimo da próxima saída
        self._pausa_ate = 0.0          # 429: ninguém sai antes disto

    def aguardar(self):
        with self._lock:
            saida = max(time.monotonic(), self._proximo, self._pausa_ate)
            self._proximo = saida + self._intervalo
        espera = saida - time.monotonic()
        if espera > 0:
            _dormir(espera)

    def bloquear(self, segundos):
        with self._lock:
            self._pausa_ate = max(self._pausa_ate, time.monotonic() + segundos)
            self._proximo   = max(self._proximo, self._pausa_ate)


_cota_compras = _CotaCompras(COTA_COMPRAS_POR_MINUTO)


def _retry_after(resp) -> float:
    """Segundos pedidos pelo servidor num 429 (+ meio segundo de folga)."""
    try:
        return float(resp.headers.get("Retry-After")) + 0.5
    except (TypeError, ValueError):
        m = re.search(r"(\d+)\s*second", resp.text or "")
        return (int(m.group(1)) if m else 15) + 0.5


def _get_compras(url, params, timeout=TIMEOUT, max_429=10):
    """GET no Compras.gov respeitando a cota. Um 429 pausa todas as threads
    pelo Retry-After e a chamada é repetida (até max_429 vezes). Erros de rede
    sobem como exceção, como num _http.get comum."""
    resp = None
    for _ in range(max_429):
        _cota_compras.aguardar()
        resp = _http.get(url, params=params, timeout=timeout)
        if resp.status_code != 429:
            return resp
        _cota_compras.bloquear(_retry_after(resp))
    return resp

# =============================================================================
# TIPOS DE BUSCA  —  espelha o seletor "tipo" do endpoint de Pesquisa de Preço
#   /modulo-pesquisa-preco/1.1_consultarMaterial_CSV?tipo={tipo}&codigo={codigo}
# =============================================================================
TIPO_CATMAT = "codigoItemCatalogo"
TIPO_PDM    = "codigoPdm"
ROTULO_TIPO = {TIPO_CATMAT: "CATMAT", TIPO_PDM: "PDM"}
TIPO_POR_ROTULO = {v: k for k, v in ROTULO_TIPO.items()}

# Detectado na 1ª requisição e reaproveitado nas demais:
#   None  → ainda não sabemos
#   True  → API aceita a assinatura nova (tipo + codigo)
#   False → API ainda na assinatura antiga (codigoItemCatalogo)
_API_ACEITA_TIPO = None

ordem_final_colunas = [
    "idCompra","idItemCompra","forma","modalidade","criterioJulgamento",
    "numeroItemCompra","descricaoItem","codigoItemCatalogo","codigoPdm","nomeUnidadeFornecimento",
    "siglaUnidadeFornecimento","nomeUnidadeMedida","capacidadeUnidadeFornecimento","siglaUnidadeMedida",
    "Unidade de Fornecimento","capacidade","quantidade","precoUnitario","Preco Total","percentualMaiorDesconto",
    "niFornecedor","nomeFornecedor","marca","codigoUasg","nomeUasg",
    "codigoMunicipio","municipio","estado","codigoOrgao","nomeOrgao",
    "poder","esfera","dataCompra","dataHoraAtualizacaoCompra","dataHoraAtualizacaoItem",
    "dataResultado","dataHoraAtualizacaoUasg","codigoClasse","nomeClasse",
]


def converter_data_para_api(data_dd_mm_yyyy: str) -> Optional[str]:
    s = data_dd_mm_yyyy.strip()
    if not s: return None
    try:
        p = s.split("-")
        if len(p) != 3: return None
        dd, mm, yyyy = p
        if len(dd) == 2 and len(mm) == 2 and len(yyyy) == 4:
            int(dd); int(mm); int(yyyy)
            return f"{yyyy}-{mm}-{dd}"
    except (ValueError, AttributeError):
        pass
    return None


def validar_e_obter_datas(ini: str, fim: str):
    i_api = f_api = None
    if ini.strip():
        i_api = converter_data_para_api(ini)
        if i_api is None:
            return None, None, f"Data de Inicio invalida: '{ini}'\nUse DD-MM-AAAA (ex: 01-01-2024)"
    if fim.strip():
        f_api = converter_data_para_api(fim)
        if f_api is None:
            return None, None, f"Data Final invalida: '{fim}'\nUse DD-MM-AAAA (ex: 31-12-2024)"
    return i_api, f_api, None


# =============================================================================
# PARSER DE PÁGINA CSV
# -----------------------------------------------------------------------------
# A API devolve CSV com campos de texto livre (descricaoItem) que podem conter
# quebras de linha cruas, ';' e aspas desbalanceadas. Três armadilhas conhecidas:
#
#   1. str.splitlines() quebra em \x0b \x0c \x1c-\x1e \x85    , que o
#      parser CSV (e o servidor) tratam como texto comum → fragmentos fantasma.
#   2. Contar ';' com str.split ignora aspas → campos citados com ';' interno
#      viram "excesso de colunas" e a linha é remontada torta.
#   3. Descartar linhas que contêm só '"' apaga o fechamento de um campo
#      multilinha → o parser engole as linhas seguintes e some com registros.
#
# A abordagem aqui é: csv.reader sobre o texto bruto (que já resolve quebras
# dentro de campos citados), remontagem no nível de CAMPO e conferência do
# número de registros contra o esperado da página.
# =============================================================================

# Quebras que str.splitlines() reconhece mas o CSV não
_QUEBRAS_FALSAS = re.compile(r"[\x0b\x0c\x1c\x1d\x1e\x85  ]")
# Caracteres de controle que o openpyxl recusa ao gravar .xlsx
_CTRL_ILEGAIS   = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
# Linhas de metadados (totalRegistros / total paginas / paginas restantes)
_RE_METADADO    = re.compile(
    r'^\s*"?\s*total\s*(de\s*)?(registros|p[áa]ginas?|p[áa]ginas?\s+restantes)\s*:',
    re.IGNORECASE)


def _limpar_campo(v):
    """Normaliza um campo: tira controles ilegais e colapsa quebras internas."""
    if not v:
        return ""
    v = v.replace("\r\n", " ").replace("\r", " ").replace("\n", " ")
    v = _QUEBRAS_FALSAS.sub(" ", v)
    v = _CTRL_ILEGAIS.sub(" ", v)
    return re.sub(r"[ \t]{2,}", " ", v).strip()


def _ler_registros(texto, quoting):
    """csv.reader sobre o texto bruto — resolve quebras dentro de campos citados."""
    try:
        return list(csv.reader(StringIO(texto, newline=""), delimiter=";",
                               quotechar='"', quoting=quoting, strict=False))
    except csv.Error:
        return []


def _remontar(registros, ncols, idx_livre):
    """
    Remonta registros quebrados.
      len < ncols → fragmento: a quebra caiu DENTRO de um campo, então o último
                    campo do fragmento e o primeiro do seguinte são as duas
                    metades do mesmo campo (por isso a junção é por campo, e
                    não pela linha inteira).
      len > ncols → ';' extra em campo sem aspas: recolhe o excedente de volta
                    para o campo de texto livre em vez de descartar a linha.
    Retorna (linhas, reparos, descartes).
    """
    linhas = []; buf = None; reparos = 0; descartes = 0
    for reg in registros:
        remontado = False
        if buf is not None:
            cabeca = (buf[-1] + " " + (reg[0] if reg else "")).strip()
            reg = buf[:-1] + [cabeca] + reg[1:]
            buf = None; remontado = True

        n = len(reg)
        if n < ncols:
            buf = reg                       # ainda incompleto — segue acumulando
            continue
        if n == ncols:
            if remontado: reparos += 1
            linhas.append(reg)
        elif idx_livre is not None and idx_livre < ncols:
            excedente = n - ncols
            reg = (reg[:idx_livre]
                   + [";".join(reg[idx_livre:idx_livre + excedente + 1])]
                   + reg[idx_livre + excedente + 1:])
            reparos += 1
            linhas.append(reg)
        else:
            descartes += 1
    if buf is not None:
        descartes += 1                      # fragmento órfão no fim da página
    return linhas, reparos, descartes


def _montar_pagina(texto, quoting):
    """Extrai (header, linhas, reparos, descartes) de uma página com um dado quoting."""
    registros = [r for r in _ler_registros(texto, quoting) if any(c.strip() for c in r)]
    registros = [r for r in registros if not _RE_METADADO.match(r[0] if r else "")]
    if not registros:
        return None, [], 0, 0
    header = [c.strip() for c in registros[0]]
    ncols  = len(header)
    if ncols < 2:
        return None, [], 0, 0
    idx_livre = header.index("descricaoItem") if "descricaoItem" in header else None
    linhas, reparos, descartes = _remontar(registros[1:], ncols, idx_livre)
    return header, linhas, reparos, descartes


def parse_pagina_csv(csv_text, esperado_na_pagina=None):
    """
    Converte o CSV de uma página em DataFrame com diagnóstico confiável.

    Retorna (df, diag), diag = {
        "linhas", "esperado", "reparos", "descartes", "modo", "ok", "motivo"
    }
    'ok' é False sempre que a página não entregou exatamente os registros
    esperados — é essa conferência (e não uma heurística de ';') que garante
    que nenhuma página problemática passe batido.
    """
    diag = {"linhas": 0, "esperado": esperado_na_pagina, "reparos": 0,
            "descartes": 0, "modo": "aspas", "ok": True, "motivo": ""}
    if not csv_text:
        diag.update(ok=False, motivo="pagina vazia")
        return pd.DataFrame(), diag

    melhor = None
    for modo, quoting in (("aspas", csv.QUOTE_MINIMAL), ("literal", csv.QUOTE_NONE)):
        header, linhas, reparos, descartes = _montar_pagina(csv_text, quoting)
        if header is None:
            continue
        cand = {"header": header, "linhas": linhas, "reparos": reparos,
                "descartes": descartes, "modo": modo}
        # Bateu o esperado no modo com aspas: não precisa da segunda leitura
        if esperado_na_pagina is not None and len(linhas) == esperado_na_pagina:
            melhor = cand
            break
        # Senão, fica com o que recupera mais registros e descarta menos
        if melhor is None or (len(linhas), -descartes) > (len(melhor["linhas"]),
                                                          -melhor["descartes"]):
            melhor = cand

    if melhor is None:
        diag.update(ok=False, motivo="cabecalho nao identificado")
        return pd.DataFrame(), diag

    dados = [[_limpar_campo(c) for c in ln] for ln in melhor["linhas"]]
    df = pd.DataFrame(dados, columns=melhor["header"], dtype=str)

    diag.update(linhas=len(df), reparos=melhor["reparos"],
                descartes=melhor["descartes"], modo=melhor["modo"])
    if esperado_na_pagina is not None and len(df) != esperado_na_pagina:
        diag["ok"] = False
        diag["motivo"] = f"{len(df)} de {esperado_na_pagina} registros"
    elif melhor["descartes"]:
        diag["ok"] = False
        diag["motivo"] = f"{melhor['descartes']} linha(s) descartada(s)"
    return df, diag


# =============================================================================
# PESQUISA DE PREÇO  —  registros de preço de um CATMAT ou de um PDM
# =============================================================================

def ler_pagina(codigo, pagina, tipo=TIPO_CATMAT,
               data_compra_inicio=None, data_compra_fim=None) -> str:
    """
    Lê uma página de Registros de Preço. Devolve o CSV da página ou uma
    mensagem que começa com "ERRO_CONEXAO" / "ERRO_REQUISICAO".

    tipo — equivale ao seletor "tipo" do endpoint de Pesquisa de Preço:
        TIPO_CATMAT ("codigoItemCatalogo") → codigo é um CATMAT
        TIPO_PDM    ("codigoPdm")          → codigo é um PDM (traz todos os
                                             CATMATs do PDM de uma só vez)

    Envia a assinatura nova (tipo + codigo). Se o servidor recusar (400/404) e a
    busca for por CATMAT, refaz com a assinatura antiga (codigoItemCatalogo),
    mantendo compatibilidade com instâncias ainda não atualizadas da API.
    """
    global _API_ACEITA_TIPO
    URL = f"{URL_BASE}/modulo-pesquisa-preco/1.1_consultarMaterial_CSV"

    base = {"tamanhoPagina": TAMANHO_PAGINA, "pagina": int(pagina)}
    if data_compra_inicio: base["dataCompraInicio"] = data_compra_inicio
    if data_compra_fim:    base["dataCompraFim"]    = data_compra_fim

    def _requisitar(params):
        """Retorna (csv_text, erro, status_http). csv_text=None quando falhou."""
        try:
            resp = _get_compras(URL, params)       # 429 já tratado pela cota
            if resp.status_code == 429:
                return None, f"ERRO_REQUISICAO: 429 persistente para {tipo} {codigo}", 429
            if resp.status_code in (400, 404):
                return None, f"ERRO_REQUISICAO: HTTP {resp.status_code}", resp.status_code
            resp.raise_for_status()
            return resp.content.decode("utf-8-sig", errors="replace"), None, 200
        except requests.exceptions.ConnectionError as e:
            return None, f"ERRO_CONEXAO: {e}", None
        except requests.exceptions.RequestException as e:
            return None, f"ERRO_REQUISICAO: {e}", None

    # ── 1ª opção: assinatura nova (tipo + codigo) ────────────────────────────
    if _API_ACEITA_TIPO is not False:
        csv_text, erro, status = _requisitar(
            dict(base, tipo=tipo, codigo=str(int(codigo))))
        if csv_text is not None:
            _API_ACEITA_TIPO = True
            return csv_text
        # Só cai para o modo legado quando o servidor recusa a assinatura
        if status not in (400, 404):
            return erro
        _API_ACEITA_TIPO = False

    # ── 2ª opção: assinatura antiga — existe apenas para CATMAT ──────────────
    if tipo != TIPO_CATMAT:
        return ("ERRO_REQUISICAO: esta instância da API não aceita busca "
                f"por {tipo}. Use CATMAT.")

    csv_text, erro, _ = _requisitar(dict(base, codigoItemCatalogo=int(codigo)))
    return csv_text if csv_text is not None else erro


def _int_do_rodape(csv_text, rotulo):
    """
    Lê um inteiro do rodapé da resposta (ex.: 'totalPaginas: 1.234').
    Tolera separador de milhar — '\\d+' sozinho capturaria apenas o '1'.
    Retorna None quando o rótulo não aparece na resposta.
    """
    m = re.search(rotulo + r"\s*:\s*([\d.,]+)", csv_text, re.IGNORECASE)
    if not m:
        return None
    digitos = re.sub(r"\D", "", m.group(1))
    return int(digitos) if digitos else None


def extrair_codigo(codigo, tipo=TIPO_CATMAT, d_ini=None, d_fim=None, pasta_corr=None):
    """
    Worker puro: busca e processa todas as páginas de um CATMAT ou de um PDM,
    conforme `tipo` (TIPO_CATMAT | TIPO_PDM).
    Pode rodar em qualquer thread — não acessa estado compartilhado.
    Retorna: (codigo, dfs_e_meta, status, reg_esp, paginas_com_perda)
      status: "ok" | "vazio" | "erro" | "cancelado"
      dfs_e_meta: lista de (df_processado, marca, num_pagina)
                  marca: "" | "reparada" | "perda"
    """
    dfs_e_meta   = []
    perdas       = []
    reg_esp      = 0
    pagina_atual = 1
    if _cancelado:
        return codigo, [], "cancelado", 0, []

    try:
        # Queda de rede na 1ª página: espera 60 s e tenta de novo, sem desistir
        # do código
        while True:
            csv_text = ler_pagina(codigo, 1, tipo, d_ini, d_fim)
            if not csv_text.startswith("ERRO_CONEXAO"):
                break
            _log(f"Sem conexão ({ROTULO_TIPO[tipo]} {codigo}) — nova tentativa em 60 s.")
            _dormir(60)
            if _cancelado:
                return codigo, [], "cancelado", 0, []
        if csv_text.startswith("ERRO_REQUISICAO"):
            return codigo, [], "erro", 0, []

        reg_esp = _int_do_rodape(csv_text, "totalRegistros") or 0
        if reg_esp == 0:
            return codigo, [], "vazio", 0, []

        # Total de páginas: o rodapé da resposta é a fonte primária. O cálculo
        # por totalRegistros entra como conferência e, quando o rodapé falta,
        # como substituto — cair em 1 nesse caso truncaria a extração nos 500
        # primeiros registros em silêncio.
        pag_rodape    = _int_do_rodape(csv_text, r"total\s*(?:de\s*)?p[áa]ginas?")
        pag_calculado = max(1, math.ceil(reg_esp / TAMANHO_PAGINA))
        total_paginas = max(pag_rodape or 0, pag_calculado)
        # Sem rodapé não há como conferir a paginação: só nesse caso vale
        # insistir enquanto as páginas voltarem cheias.
        confiar_no_rodape = pag_rodape is not None

        while True:
            # Cancelamento também ENTRE PÁGINAS de um mesmo código
            if _cancelado:
                break

            esperado_pag = max(0, min(TAMANHO_PAGINA,
                                      reg_esp - TAMANHO_PAGINA * (pagina_atual - 1)))
            df_pag, diag = parse_pagina_csv(csv_text, esperado_pag or None)
            # "perda"    → não foi possível reconstituir a página fielmente
            # "reparada" → houve conserto, mas todos os registros foram recuperados
            if not diag["ok"]:
                marca = "perda"
            elif diag["reparos"]:
                marca = "reparada"
            else:
                marca = ""

            if marca == "perda":
                perdas.append(str(pagina_atual))
                if pasta_corr:
                    dest = os.path.join(pasta_corr,
                        f"cod_{codigo}_pag_{pagina_atual}_corr.csv")
                    try:
                        with open(dest, "w", encoding="utf-8-sig") as f:
                            f.write(csv_text)
                    except Exception:
                        pass

            if df_pag is not None and not df_pag.empty:
                if tipo == TIPO_PDM:
                    # Busca por PDM devolve vários CATMATs: preserva o
                    # codigoItemCatalogo original e apenas anota o PDM de origem
                    df_pag.loc[:, "codigoPdm"] = str(codigo)
                else:
                    df_pag.loc[:, "codigoItemCatalogo"] = str(codigo)
                df_proc = processar_dataframe_final(df_pag, ordem_final_colunas)
                dfs_e_meta.append((df_proc, marca, pagina_atual))

            # Com rodapé, ele manda: nenhuma requisição além do que ele informa.
            # Sem rodapé, página cheia é o único indício de que ainda há dados.
            insistir = (not confiar_no_rodape
                        and df_pag is not None and len(df_pag) >= TAMANHO_PAGINA)
            pagina_atual += 1
            if pagina_atual > total_paginas and not insistir:
                break
            if pagina_atual > _MAX_PAGINAS:      # trava de segurança
                break
            # Sem pausa fixa entre páginas: o ritmo é o da cota (_get_compras)
            csv_text = ler_pagina(codigo, pagina_atual, tipo, d_ini, d_fim)
            if csv_text.startswith("ERRO_"):
                # O relatório de integridade acusa a divergência deste código
                _log(f"{ROTULO_TIPO[tipo]} {codigo} pág {pagina_atual}: {csv_text[:160]}"
                     " — páginas seguintes não baixadas.")
                break

        return codigo, dfs_e_meta, ("ok" if dfs_e_meta else "vazio"), reg_esp, perdas

    except Exception:
        return codigo, [], "erro", 0, []


def processar_dataframe_final(df: pd.DataFrame, ordem_colunas: List[str]) -> pd.DataFrame:
    if df.empty: return df
    fc = df.columns[0]
    df = df[~df[fc].astype(str).str.contains("totalRegistros|totalPaginas",
                                              case=False, na=False)].copy()
    if df.empty: return df

    # Mapa sigla → nome completo para preencher nomeUnidadeFornecimento ausente
    _SIGLA_NOME = {
        "FR-AM": "Frasco-Ampola", "FR": "Frasco", "CAPS": "Cápsula",
        "COMPR": "Comprimido", "AM": "Ampola", "UN": "Unidade",
        "SER": "Seringa", "BIS": "Bisnaga", "BLIS": "Blister",
        "BOL": "Bolsa", "BOM": "Bombona", "CA": "Cartucho",
        "CI": "Curie", "CJ": "Conjunto", "DOSE(S)": "Dose(s)",
        "DOSES": "Dose(s)", "DRAG": "Drágea", "EMB": "Embalagem",
        "EMP": "Emplastro", "ENV": "Envelope", "FLAC": "Flaconete",
        "G": "Grama", "GL": "Galão", "GLOB": "Glóbulo",
        "KG": "Quilograma", "L": "Litro", "MCG": "Micrograma",
        "MCU": "Milicurie", "MG": "Miligrama",
        "MIL CTE": "Milheiro de Cartelas", "ML": "Mililitro",
        "PAST": "Pastilha", "POTE": "Pote", "RO": "Rolo",
        "SAC": "Sachê", "SUP": "Supositório", "TAB": "Tablete",
        "TBO": "Tubo", "TBTE": "Tubete", "UI": "Unid. Internacional",
    }

    def _val(row, col):
        v = row.get(col)
        s = str(v).strip() if pd.notna(v) else ""
        return "" if s in ("", "nan", "None", "null") else s

    def uf(row):
        nome  = _val(row, "nomeUnidadeFornecimento")
        sigla = _val(row, "siglaUnidadeFornecimento")
        cap   = _val(row, "capacidadeUnidadeFornecimento")
        medi  = _val(row, "siglaUnidadeMedida")

        # Se nomeUnidade vazio, tentar preencher pela sigla
        if not nome and sigla:
            nome = _SIGLA_NOME.get(sigla.upper(), sigla)

        # capacidade: ignorar se for 0,00 ou 0.00 ou 0
        try:
            cap_num = float(cap.replace(".", "").replace(",", ".")) if cap else 0
        except Exception:
            cap_num = 0
        cap_valida = bool(cap) and cap_num != 0

        # Montar: Nome + capacidade + siglaUnidadeMedida
        # Se não houver capacidade válida, ignorar também siglaUnidadeMedida
        partes = [nome] if nome else []
        if cap_valida:
            partes.append(cap)
            if medi:
                partes.append(medi)
        return " ".join(partes)

    df["Unidade de Fornecimento"] = df.apply(uf, axis=1)

    def tof(v):
        if pd.isna(v): return 0.0
        try: return float(str(v).replace(".", "").replace(",", "."))
        except: return 0.0

    # Sem os .get() a ausência de uma coluna vira KeyError, que o worker
    # captura no except genérico e reporta como "erro de API" — uma mudança de
    # schema na origem apareceria como indisponibilidade, sem pista no log.
    if "precoUnitario" in df.columns and "quantidade" in df.columns:
        df["Preco Total"] = df["precoUnitario"].apply(tof) * df["quantidade"].apply(tof)
    else:
        df["Preco Total"] = 0.0
    for col in ["nomeUnidadeMedida","percentualMaiorDesconto"]:
        if col in df.columns and (df[col].isnull().all() or
                                   df[col].astype(str).str.strip().eq("").all()):
            df = df.drop(columns=[col])
    exist = [c for c in ordem_colunas if c in df.columns]
    extra = [c for c in df.columns if c not in exist]
    return df[exist + extra]


# =============================================================================
# CATÁLOGO DE MATERIAIS  —  Classe → PDMs → CATMATs
# =============================================================================

def _normalizar_campo(item: dict, *candidatos, default=""):
    """Retorna o primeiro campo encontrado no dict entre os candidatos."""
    for c in candidatos:
        if c in item and item[c] is not None:
            return item[c]
    return default


def _concluidos(futuros):
    """as_completed que não segura o Ctrl+C: espera em fatias de 0,5 s e para
    de entregar resultados quando a extração é cancelada."""
    pendentes = set(futuros)
    while pendentes and not _cancelado:
        prontos, pendentes = wait(pendentes, timeout=0.5, return_when=FIRST_COMPLETED)
        yield from prontos


def buscar_pdms_por_classe(codigo_classe: int, max_tentativas: int = 3) -> Optional[List[int]]:
    """Códigos de todos os PDMs de uma classe, com retry automático e backoff.
    None se a API falhar (ou a classe não tiver PDMs)."""
    URL = f"{URL_BASE}/modulo-material/3_consultarPdmMaterial"
    todos = []; pagina_atual = 1; total_paginas = 1

    while pagina_atual <= total_paginas:
        data = None
        for tentativa in range(max_tentativas):
            if _cancelado:
                return None
            try:
                # 429 é tratado pela cota (_get_compras): espera o Retry-After
                resp = _get_compras(URL, params={
                    "codigoClasse": codigo_classe, "pagina": pagina_atual,
                    "tamanhoPagina": 500, "bps": "false"
                }, timeout=TIMEOUT)
                resp.raise_for_status()
                data = resp.json()
                break
            except requests.exceptions.ConnectionError as e:
                espera = 3 * (tentativa + 1)
                _log(f"Erro de conexão na classe {codigo_classe} pág {pagina_atual}: {e} "
                     f"— aguardando {espera} s (tentativa {tentativa+1})")
                _dormir(espera)
            except Exception as e:
                espera = 2 * (tentativa + 1)
                _log(f"Erro na classe {codigo_classe} pág {pagina_atual}: {e} "
                     f"— aguardando {espera} s (tentativa {tentativa+1})")
                _dormir(espera)

        if data is None:
            _log(f"Falha definitiva: classe {codigo_classe} pág {pagina_atual} "
                 f"após {max_tentativas} tentativas")
            return None

        todos.extend(data.get("resultado") or [])
        if pagina_atual == 1:
            total_registros = int(data.get("totalRegistros", 0))
            total_paginas = (math.ceil(total_registros / 500)
                             if total_registros > 0 else 1)
        pagina_atual += 1       # o ritmo entre páginas é dado pela cota

    # A API pode retornar nomes de campo variados
    pdms = []
    for item in todos:
        try:
            pdms.append(int(_normalizar_campo(item, "codigoPdm", "codigo", "id", "codigoItem")))
        except (TypeError, ValueError):
            pass
    return list(dict.fromkeys(pdms)) or None


def buscar_catmats_por_pdm(codigos_pdm, max_workers=5):
    """
    Busca os CATMATs de vários PDMs em paralelo. As threads cobrem a latência;
    o ritmo real é o da cota (_get_compras). Devolve (catmats, pdms_com_erro).
    """
    URL   = f"{URL_BASE}/modulo-material/4_consultarItemMaterial"
    total = len(codigos_pdm)

    def _fetch_pdm(pdm_code):
        """Worker: todas as páginas de CATMATs de um PDM. None = erro."""
        itens = []; pagina_atual = 1; total_paginas = 1
        try:
            while pagina_atual <= total_paginas:
                if _cancelado:
                    return None
                # 429 é tratado pela cota; se persistir, raise_for_status
                # marca o PDM como erro e ele entra na próxima tentativa
                resp = _get_compras(URL, params={
                    "codigoPdm": pdm_code, "pagina": pagina_atual,
                    "tamanhoPagina": 500, "bps": "false"
                }, timeout=TIMEOUT)
                resp.raise_for_status()
                data = resp.json()
                itens.extend(data.get("resultado") or [])
                if pagina_atual == 1:
                    total_paginas = int(data.get("totalPaginas") or 1)
                pagina_atual += 1
        except Exception:
            return None
        return itens

    catmats, pdms_com_erro, feitos = [], [], 0
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futuros = {executor.submit(_fetch_pdm, pdm): pdm for pdm in codigos_pdm}
        for fut in _concluidos(futuros):
            itens = fut.result()
            if itens is None:
                pdms_com_erro.append(futuros[fut])
            else:
                for item in itens:
                    try:
                        catmats.append(int(item["codigoItem"]))
                    except (KeyError, TypeError, ValueError):
                        pass
            feitos += 1
            if feitos % 10 == 0 or feitos == total:
                _log(f"  PDMs consultados: {feitos}/{total}")
        if _cancelado:
            executor.shutdown(cancel_futures=True)
    return catmats, pdms_com_erro


def catmats_dos_pdms(pdms) -> List[int]:
    """CATMATs de uma lista de PDMs, com até 3 tentativas para os PDMs com erro."""
    catmats, pendentes = [], list(pdms)
    for tentativa, espera in enumerate((0, 3, 8), 1):
        if not pendentes or _cancelado:
            break
        if espera:
            _log(f"{tentativa}ª tentativa para {len(pendentes)} PDMs com erro "
                 f"(aguardando {espera} s)...")
            _dormir(espera)
        achados, pendentes = buscar_catmats_por_pdm(pendentes)
        catmats += achados
    if pendentes and not _cancelado:
        _log("PDMs sem resposta após 3 tentativas: " + ", ".join(map(str, pendentes)))
    return list(dict.fromkeys(catmats))


# =============================================================================
# DESCRIÇÃO DO BPS  —  troca o descricaoItem do Compras.gov pelo texto do BPS
# -----------------------------------------------------------------------------
# GET /economia-da-saude/bps?codigoCatmat=...&pagina=1&tamanhoPagina=1
# API de Dados Abertos do Ministério da Saúde: outro servidor, fora da cota do
# Compras.gov. Uma chamada leve por CATMAT, sem filtro de data — basta um
# registro de qualquer compra para ter a descrição. CATMAT nunca comprado no
# BPS fica com o texto do Compras.gov. É o mesmo caminho do ExtratorCatmat com
# "Usar a descrição do BPS" marcado e a fonte BPS desmarcada.
# =============================================================================

URL_BPS = "https://apidadosabertos.saude.gov.br/economia-da-saude/bps"
_pool_bps = ThreadPoolExecutor(max_workers=4, thread_name_prefix="bps")
# Descrição do BPS por CATMAT ("" = CATMAT sem nenhuma compra no BPS)
_cache_desc_bps: dict = {}
_cache_desc_lock = threading.Lock()

# ── Limpeza da descrição: ESPELHO do bloco ds_catmat do extracao_bps.sql ────
# Mesmas regras, na mesma ordem, para que a descrição que o extrator grava
# seja idêntica à da extração SQL. Mudou uma regra lá? Mude aqui também (e no
# ExtratorCatmat.py).
_TAGS_BPS = re.compile(r"</?(div|span|p|br|b|i|strong|em)[^>]*>", re.IGNORECASE)
_ENTIDADES_BPS = (
    ("&#193;", "Á"), ("&#194;", "Â"), ("&#195;", "Ã"), ("&#199;", "Ç"),
    ("&#201;", "É"), ("&#202;", "Ê"), ("&#205;", "Í"), ("&#211;", "Ó"),
    ("&#212;", "Ô"), ("&#213;", "Õ"), ("&#218;", "Ú"), ("&#220;", "Ü"),
    ("&#225;", "á"), ("&#226;", "â"), ("&#227;", "ã"), ("&#231;", "ç"),
    ("&#233;", "é"), ("&#234;", "ê"), ("&#237;", "í"), ("&#243;", "ó"),
    ("&#244;", "ô"), ("&#245;", "õ"), ("&#250;", "ú"), ("&#252;", "ü"),
    ("&#186;", "º"), ("&#170;", "ª"),
)
_ESP = r"[ \t\n\r\f\v]"            # o [[:space:]] do PostgreSQL
_RE_DOIS_PONTOS_BPS = re.compile(_ESP + r"*:" + _ESP + r"*")
_RE_ESP_VIRGULA_BPS = re.compile(_ESP + r"+,")
_RE_ESPACOS_BPS     = re.compile(_ESP + r"+")


def limpar_descricao_bps(texto) -> str:
    """Descrição do BPS limpa exatamente como o extracao_bps.sql limpa."""
    if not texto:
        return ""
    s = _TAGS_BPS.sub("", str(texto))
    for entidade, caractere in _ENTIDADES_BPS:
        s = s.replace(entidade, caractere)
    s = s.replace("¿", "°").replace("*", "")
    s = _RE_DOIS_PONTOS_BPS.sub(": ", s)
    s = _RE_ESP_VIRGULA_BPS.sub(",", s)
    s = _RE_ESPACOS_BPS.sub(" ", s)
    return s.strip(" ")


def _cod_bps(valor) -> str:
    """CATMAT como chave: só dígitos, sem zeros à esquerda."""
    s = str(valor if valor is not None else "").strip()
    return (s.lstrip("0") or "0") if s.isdigit() else ""


def _get_bps(params, tentativas=4):
    """Uma página da API do BPS. Devolve (lista, erro); lista=None se falhou."""
    ultimo = ""
    for t in range(tentativas):
        if _cancelado:
            return None, "cancelado"
        try:
            r = _http.get(URL_BPS, params=params, timeout=TIMEOUT)
            if r.status_code == 429 or r.status_code >= 500:
                ultimo = f"HTTP {r.status_code}"
                _dormir(15 * (t + 1) if r.status_code == 429 else 3 * (t + 1))
                continue
            r.raise_for_status()
            return r.json().get("bps") or [], None
        except (requests.exceptions.RequestException, ValueError, AttributeError) as e:
            ultimo = f"{type(e).__name__}: {e}"
            _dormir(3 * (t + 1))
    return None, ultimo


def _descricao_bps(codigo):
    """Descrição de um CATMAT no BPS: basta um registro de qualquer data.
    None = o BPS não respondeu; "" = CATMAT nunca comprado no BPS."""
    lote, _ = _get_bps({"codigoCatmat": _cod_bps(codigo), "pagina": 1,
                        "tamanhoPagina": 1})
    if lote is None:
        return None
    return limpar_descricao_bps(lote[0].get("descricaoItem")) if lote else ""


def aplicar_descricao_bps(dfs_e_meta) -> dict:
    """Troca, nas páginas de um código, o descricaoItem do Compras.gov pela
    descrição do BPS do mesmo CATMAT (coluna codigoItemCatalogo — na busca por
    PDM cada registro traz o seu). Cada CATMAT é consultado uma vez por
    execução; falha não entra no cache e é tentada de novo se o CATMAT voltar.

    Devolve {"trocadas": linhas com o texto trocado,
             "sem_bps": CATMATs nunca comprados no BPS (texto do Compras.gov),
             "falhas":  CATMATs que o BPS não respondeu (texto do Compras.gov)}
    """
    usados = set()
    for df, _m, _p in dfs_e_meta:
        if "codigoItemCatalogo" in df.columns:
            usados.update(_cod_bps(c) for c in df["codigoItemCatalogo"])
    usados.discard("")
    with _cache_desc_lock:
        faltam = [c for c in usados if c not in _cache_desc_bps]
    futuros = [(c, _pool_bps.submit(_descricao_bps, c)) for c in faltam]
    falhas = []
    for c, fut in futuros:
        desc = fut.result()
        if desc is None:
            falhas.append(c)
            continue
        with _cache_desc_lock:
            _cache_desc_bps[c] = desc
    with _cache_desc_lock:
        mapa = {c: _cache_desc_bps.get(c, "") for c in usados}
    sem_bps = sorted((c for c, d in mapa.items() if not d and c not in falhas), key=int)
    mapa = {c: d for c, d in mapa.items() if d}

    trocadas = 0
    for df, _m, _p in dfs_e_meta:
        if not mapa or "codigoItemCatalogo" not in df.columns \
                or "descricaoItem" not in df.columns:
            continue
        novo = df["codigoItemCatalogo"].map(_cod_bps).map(mapa)
        trocar = novo.notna() & (novo != df["descricaoItem"])
        if trocar.any():
            df.loc[trocar, "descricaoItem"] = novo[trocar]
            trocadas += int(trocar.sum())
    return {"trocadas": trocadas, "sem_bps": sem_bps, "falhas": sorted(falhas, key=int)}


def extrair_codigo_bps(codigo, tipo=TIPO_CATMAT, d_ini=None, d_fim=None,
                       pasta_corr=None, descricao_bps=True):
    """extrair_codigo + a troca da descrição pelo texto do BPS, na thread do
    worker — a consulta ao BPS corre em paralelo com os outros códigos.
    Devolve a 5-tupla de extrair_codigo + o resumo do BPS (None quando a
    troca está desligada ou o código não trouxe registros)."""
    res = extrair_codigo(codigo, tipo, d_ini, d_fim, pasta_corr)
    if not descricao_bps or res[2] != "ok":
        return res + (None,)
    try:
        bps = aplicar_descricao_bps(res[1])
    except Exception as e:                  # BPS nunca derruba a extração do Compras.gov
        _log(f"{ROTULO_TIPO[tipo]} {codigo}: descrição do BPS não aplicada "
             f"({type(e).__name__}: {e}) — ficou o texto do Compras.gov.")
        bps = {"trocadas": 0, "sem_bps": [], "falhas": ["(erro interno)"]}
    return res + (bps,)


# =============================================================================
# GRAVAÇÃO  —  .xlsx em streaming ou .csv em append, com rollover a cada 1 mi
# =============================================================================

def _nome_da_parte(base_filename, ext_padrao, parte):
    """base.xlsx -> base_part1.xlsx, base_part2.xlsx ..."""
    base, ext = os.path.splitext(base_filename)
    if not ext or ext.lower() != ext_padrao: ext = ext_padrao
    return f"{base}_part{parte}{ext}"


class ExcelChunkWriter:
    """
    Escreve .xlsx em modo STREAMING (openpyxl write_only).

    O modo padrão do openpyxl mantém a planilha inteira em memória e só serializa
    tudo no save(), o que concentra o custo no encerramento. Em write_only as
    linhas vão para disco conforme chegam: numa extração de 230 mil registros
    isso troca ~70 s de espera no final por ~7 s, e ~2,5 GB de RAM por nada.

    Contrapartida: em write_only o workbook só pode ser salvo UMA vez, então não
    há como reescrever o arquivo periodicamente. A proteção contra queda no meio
    da execução é um espelho .parcial.csv, gravado em append e apagado quando o
    .xlsx é fechado com sucesso.
    """

    def __init__(self, base_filename, sheet_name="Dados CATMAT",
                 max_rows_per_file=1_000_000):
        self.base_filename = base_filename
        self.sheet_name    = sheet_name
        self.max_rows      = max_rows_per_file
        self.part          = 1
        self.files_saved: List[str] = []
        self.header: List[str] = []
        self._linhas       = 0
        self._finalizado   = False
        self._esp_f = self._esp_w = None
        self._new_workbook()

    def _filepath(self):
        return _nome_da_parte(self.base_filename, ".xlsx", self.part)

    def _espelho_path(self):
        return os.path.splitext(self.base_filename)[0] + ".parcial.csv"

    def _new_workbook(self):
        # write_only exige create_sheet(); wb.active não existe nesse modo
        self.wb = Workbook(write_only=True)
        self.ws = self.wb.create_sheet(self.sheet_name)
        self._linhas = 0
        if self.header:
            self.ws.append(self.header)

    def _abrir_espelho(self):
        """Espelho .parcial.csv — rede de segurança enquanto o .xlsx não fecha."""
        try:
            self._esp_f = open(self._espelho_path(), "w", encoding="utf-8-sig", newline="")
            self._esp_w = csv.writer(self._esp_f, delimiter=";")
            self._esp_w.writerow(self.header)
        except Exception:
            self._esp_f = self._esp_w = None     # sem espelho é melhor que falhar

    def _fechar_espelho(self, apagar):
        if self._esp_f is None:
            return
        try:
            self._esp_f.close()
        except Exception:
            pass
        if apagar:
            try:
                os.remove(self._espelho_path())
            except Exception:
                pass
        self._esp_f = self._esp_w = None

    def _rollover(self):
        path = self._filepath(); self.wb.save(path); self.files_saved.append(path)
        self.part += 1; self._new_workbook()

    def write_dataframe(self, df: pd.DataFrame):
        if df is None or df.empty: return
        if not self.header:
            self.header = list(df.columns)
            self.ws.append(self.header)
            self._abrir_espelho()
        # O conjunto de colunas varia entre páginas (processar_dataframe_final
        # descarta colunas 100% vazias): alinha pelo cabeçalho da 1ª página
        df = df.reindex(columns=self.header)
        for linha in df.itertuples(index=False, name=None):
            if self._linhas >= self.max_rows:
                self._rollover()
            # openpyxl levanta IllegalCharacterError em caracteres de controle,
            # frequentes no texto livre vindo da API — sanitiza na gravação
            limpa = [None if pd.isna(v) else
                     (_CTRL_ILEGAIS.sub(" ", v) if isinstance(v, str) else v)
                     for v in linha]
            self.ws.append(limpa)
            if self._esp_w is not None:
                self._esp_w.writerow(["" if v is None else v for v in limpa])
            self._linhas += 1
        if self._esp_f is not None:
            try:
                self._esp_f.flush()
            except Exception:
                pass

    def _descartar_workbook(self):
        """
        Um workbook write_only coletado sem save() deixa os geradores internos do
        openpyxl abertos, e o lxml despeja 'Exception ignored ... LxmlSyntaxError'
        no stderr durante o garbage collector. close() encerra os streams.
        """
        try:
            self.ws.close()
        except Exception:
            pass
        try:
            self.wb.close()
        except Exception:
            pass

    def finalize(self) -> List[str]:
        if self._finalizado:
            return self.files_saved
        self._finalizado = True
        ok = True
        if self._linhas > 0:
            path = self._filepath()
            try:
                self.wb.save(path)
                if path not in self.files_saved: self.files_saved.append(path)
            except Exception as e:
                ok = False          # mantém o espelho: é tudo o que restou
                _log(f"[ERRO] {path} não foi salvo ({e}). "
                     f"Os dados estão em {self._espelho_path()}")
        else:
            self._descartar_workbook()   # nada a salvar: fecha sem ruído
        # Espelho só é descartado quando o .xlsx foi fechado com sucesso
        self._fechar_espelho(apagar=ok)
        return self.files_saved


class CSVChunkWriter:
    def __init__(self, base_filename, sep=";", encoding="utf-8-sig",
                 max_rows_per_file=1_000_000):
        self.base_filename = base_filename; self.sep = sep
        self.encoding = encoding; self.max_rows = max_rows_per_file
        self.part = 1; self.current_row_count = 0
        self.files_saved: List[str] = []; self.header_written = False
        self.header: List[str] = []

    def _filepath(self):
        return _nome_da_parte(self.base_filename, ".csv", self.part)

    def write_dataframe(self, df: pd.DataFrame):
        if df is None or df.empty: return
        # O conjunto de colunas varia entre páginas (processar_dataframe_final
        # descarta colunas 100% vazias). Sem reindexar pelo cabeçalho da 1ª
        # página, o append gravaria valores sob colunas erradas.
        if not self.header:
            self.header = list(df.columns)
        df = df.reindex(columns=self.header)
        if self.current_row_count + len(df) > self.max_rows:
            self.part += 1; self.current_row_count = 0; self.header_written = False
        path = self._filepath()
        df.to_csv(path, sep=self.sep, index=False,
                  mode="a" if self.header_written else "w",
                  header=not self.header_written, encoding=self.encoding)
        self.header_written = True; self.current_row_count += len(df)
        if path not in self.files_saved: self.files_saved.append(path)

    def finalize(self) -> List[str]:
        return self.files_saved


class Destino:
    """Para onde vai cada página: um arquivo só, ou um arquivo por classe.

    A classe de um código vem do mapa (coluna "classe" do arquivo de entrada)
    quando há um; sem ele, do campo codigoClasse de cada registro, que a API já
    devolve — separar por classe não custa nenhuma requisição a mais.
    """

    def __init__(self, pasta, fmt, base="dados_completos_extraidos",
                 por_classe=False, mapa=None):
        self.pasta, self.fmt, self.base = pasta, fmt, base
        self.por_classe    = por_classe
        self.classe_do_cod = {c: cl for cl, cods in (mapa or {}).items() for c in cods}
        self.writers: dict = {}              # nome base do arquivo -> writer

    def _writer(self, base):
        w = self.writers.get(base)
        if w is None:
            caminho = os.path.join(self.pasta, base)
            w = self.writers[base] = (CSVChunkWriter(caminho + ".csv") if self.fmt == "csv"
                                      else ExcelChunkWriter(caminho + ".xlsx"))
        return w

    @staticmethod
    def _base_classe(classe):
        classe = re.sub(r'[\\/:*?"<>|]', "_", str(classe or "")).strip()
        return f"classe_{classe or 'sem_classe'}"

    def gravar(self, codigo, df):
        if not self.por_classe:
            self._writer(self.base).write_dataframe(df)
        elif self.classe_do_cod or "codigoClasse" not in df.columns:
            self._writer(self._base_classe(self.classe_do_cod.get(codigo))).write_dataframe(df)
        else:
            # Classe lida do próprio registro: uma partição por classe encontrada
            chaves = (df["codigoClasse"].astype(str).str.strip()
                      .replace({"nan": "", "None": ""}))
            for classe, parte in df.groupby(chaves, sort=False):
                self._writer(self._base_classe(classe)).write_dataframe(parte)

    def finalizar(self) -> List[str]:
        return [p for w in self.writers.values() for p in w.finalize()]


# =============================================================================
# EXTRAÇÃO  +  RELATÓRIO DE INTEGRIDADE
# =============================================================================

class Resultado:
    """Contagens de uma extração — a base do relatório de integridade."""

    def __init__(self, codigos, tipo):
        self.codigos     = list(dict.fromkeys(codigos))
        self.tipo        = tipo
        self.esperados   = {}       # código -> totalRegistros informado pela API
        self.baixados    = {}       # código -> registros efetivamente gravados
        self.perdas      = {}       # código -> páginas com perda
        self.reparadas   = 0        # páginas consertadas sem perder registro
        self.vazios      = 0
        self.processados = set()
        self.erros       = set()    # erro de API ainda sem sucesso
        self.descricao_bps = False
        self.bps         = {}       # código -> resumo de aplicar_descricao_bps

    @property
    def total(self):
        return sum(self.baixados.values())

    def registrar(self, codigo, dfs_e_meta, reg_esp, perdas, bps=None):
        self.processados.add(codigo); self.erros.discard(codigo)
        self.esperados[codigo] = reg_esp
        self.baixados[codigo]  = sum(len(df) for df, _, _ in dfs_e_meta)
        if perdas:
            self.perdas[codigo] = perdas
        self.reparadas += sum(1 for _, marca, _ in dfs_e_meta if marca == "reparada")
        if not dfs_e_meta:
            self.vazios += 1
        if bps is not None:
            self.bps[codigo] = bps

    def status_bps(self, codigo) -> str:
        """Coluna 'descricao BPS' do relatório de integridade."""
        b = self.bps.get(codigo)
        if b is None:
            return ""
        partes = []
        if b["falhas"]:
            partes.append("ERRO: BPS sem resposta p/ " + ", ".join(b["falhas"]))
        if b["sem_bps"]:
            partes.append("sem compra no BPS: " + ", ".join(b["sem_bps"]))
        return " | ".join(partes) + " (texto do Compras.gov)" if partes else "OK"


def extrair(codigos, tipo, d_ini, d_fim, destino, pasta_corr=None,
            descricao_bps=True) -> Resultado:
    """
    Extrai os Registros de Preço de uma lista de códigos (CATMATs ou PDMs) e
    grava cada página em `destino`. Paralelo até WORKERS_COMPRAS: quem limita o
    ritmo é a cota. Códigos com erro de API voltam numa fila de retry
    (15 s → 30 s). A gravação acontece só nesta thread.
    descricao_bps: troca o descricaoItem pelo texto do BPS antes de gravar.
    """
    res       = Resultado(codigos, tipo)
    res.descricao_bps = descricao_bps
    rotulo    = ROTULO_TIPO[tipo]
    n         = len(res.codigos)
    pendentes = list(res.codigos)
    for espera in (0, 15, 30):
        if not pendentes or _cancelado:
            break
        if espera:
            _log(f"Retry de {len(pendentes)} {rotulo}(s) com erro (aguardando {espera} s)...")
            _dormir(espera)
        erros = []
        with ThreadPoolExecutor(max_workers=WORKERS_COMPRAS) as executor:
            futuros = [executor.submit(extrair_codigo_bps, c, tipo, d_ini, d_fim,
                                       pasta_corr, descricao_bps)
                       for c in pendentes]
            for fut in _concluidos(futuros):
                codigo, dfs_e_meta, status, reg_esp, perdas, bps = fut.result()
                if status == "cancelado":
                    continue
                if status == "erro":
                    erros.append(codigo); res.erros.add(codigo)
                    _log(f"{rotulo} {codigo}: erro na API.")
                    continue
                for df_proc, _, _ in dfs_e_meta:
                    destino.gravar(codigo, df_proc)
                res.registrar(codigo, dfs_e_meta, reg_esp, perdas, bps)

                bx  = res.baixados[codigo]
                txt = f"[{len(res.processados)}/{n}] {rotulo} {codigo}: {_milhar(bx)} registros"
                if bx != reg_esp:
                    txt += f" de {_milhar(reg_esp)} esperados"
                if bps is not None:
                    txt += f" | descrição do BPS em {_milhar(bps['trocadas'])} linha(s)"
                _log(txt + ".")
                for _, marca, pag in dfs_e_meta:
                    if marca == "reparada":
                        _log(f"    pág {pag}: reparada (íntegra).")
                for pag in perdas:
                    _log(f"    pág {pag}: registros perdidos.")
                if bps is not None and bps["falhas"]:
                    _log("    BPS sem resposta para o(s) CATMAT(s) "
                         + ", ".join(bps["falhas"]) + " — ficou o texto do Compras.gov.")
            if _cancelado:
                _log("Cancelando: aguardando as requisições em andamento...")
                executor.shutdown(cancel_futures=True)
        pendentes = erros
    if pendentes and not _cancelado:
        _log(f"{len(pendentes)} {rotulo}(s) sem resposta após 3 tentativas: "
             + ", ".join(map(str, pendentes)))
    return res


def gravar_relatorio(caminho, res: Resultado, aba="Relatorio Integridade"):
    """Uma linha por código: esperados x baixados, páginas com perda e status
    (+ a origem da descrição, quando a troca pelo texto do BPS está ligada)."""
    wb = Workbook(); ws = wb.active; ws.title = aba[:31]
    ws.append([res.tipo, "esperados", "baixados", "paginas", "status"]
              + (["descricao BPS"] if res.descricao_bps else []))
    for c in res.codigos:
        bx = int(res.baixados.get(c, 0))
        ex = int(res.esperados.get(c, 0))
        if c in res.processados:
            d  = abs(ex - bx)
            st = ("OK" if d == 0 else
                  f"OK (divergencia: {bx}/{ex})" if d <= 2 else
                  f"Inconsistencia Grave ({bx}/{ex})")
        elif c in res.erros:
            st = "ERRO_API_PERSISTENTE"
        else:
            st = "NAO PROCESSADO (cancelado)"
        ws.append([c, ex, bx, ", ".join(res.perdas.get(c, [])), st]
                  + ([res.status_bps(c)] if res.descricao_bps else []))
    wb.save(caminho)


# =============================================================================
# FLUXOS
# =============================================================================

def fluxo_codigos(codigos, tipo, d_ini, d_fim, pasta, fmt,
                  por_classe=False, mapa=None, pasta_corr=None, descricao_bps=True):
    """Lista de CATMATs/PDMs → um arquivo (ou um por classe) + relatório."""
    destino = Destino(pasta, fmt, por_classe=por_classe, mapa=mapa)
    try:
        res = extrair(codigos, tipo, d_ini, d_fim, destino, pasta_corr, descricao_bps)
    finally:
        arquivos = destino.finalizar()      # salva o que houver, mesmo cancelado
    rel = os.path.join(pasta, "Relatorio_Integridade.xlsx")
    gravar_relatorio(rel, res)
    return [res], arquivos + [rel]


def fluxo_classes(classes, tipo, d_ini, d_fim, pasta, fmt, pasta_corr=None,
                  descricao_bps=True):
    """
    Para cada classe: PDMs → CATMATs → Registros de Preço → classe_XXXX +
    Relatorio_Integridade_XXXX. Com tipo=PDM a expansão PDM → CATMAT é
    dispensada: a própria Pesquisa de Preço devolve todos os itens do PDM.
    Classes sem PDMs são reprocessadas depois das demais (3 s → 8 s).
    """
    resultados, arquivos = [], []

    def _processar_classe(classe, idx, total):
        """False = a busca de PDMs falhou (a classe volta para a fila)."""
        _log()
        _log("-" * 50)
        _log(f"CLASSE {classe}  ({idx}/{total})")
        _log("-" * 50)
        pdms = buscar_pdms_por_classe(int(classe))
        if pdms is None:
            return False
        _log(f"Classe {classe}: {len(pdms)} PDMs.")

        if tipo == TIPO_PDM:
            _log(f"Classe {classe}: extração direta dos {len(pdms)} PDMs (sem expandir CATMATs).")
            codigos = pdms
        else:
            codigos = catmats_dos_pdms(pdms)
            if _cancelado:
                return True
            if not codigos:
                _log(f"Classe {classe}: nenhum CATMAT. Pulando.")
                return True
            _log(f"Classe {classe}: {len(codigos)} CATMATs.")

        destino = Destino(pasta, fmt, base=f"classe_{classe}")
        try:
            res = extrair(codigos, tipo, d_ini, d_fim, destino, pasta_corr, descricao_bps)
        finally:
            gerados = destino.finalizar()
        rel = os.path.join(pasta, f"Relatorio_Integridade_{classe}.xlsx")
        gravar_relatorio(rel, res, aba=f"Integridade_{classe}")
        nomes = ", ".join(os.path.basename(p) for p in gerados) or "(sem dados)"
        _log(f"Classe {classe}: {_milhar(res.total)} registros -> {nomes}")
        resultados.append(res)
        arquivos.extend(gerados + [rel])
        return True

    falhas = []
    for idx, classe in enumerate(classes, 1):
        if _cancelado:
            break
        if not _processar_classe(classe, idx, len(classes)) and not _cancelado:
            _log(f"Classe {classe}: sem PDMs na 1ª tentativa. Fila de retry.")
            falhas.append(classe)

    for tentativa, espera in enumerate((3, 8), 2):
        if not falhas or _cancelado:
            break
        _log()
        _log(f"{len(falhas)} classe(s) sem PDMs — tentativa {tentativa}/3 "
             f"(aguardando {espera} s)...")
        _dormir(espera)
        ainda = []
        for idx, classe in enumerate(falhas, 1):
            if _cancelado:
                break
            if not _processar_classe(classe, idx, len(falhas)):
                ainda.append(classe)
        falhas = ainda

    if falhas and not _cancelado:
        _log("Classes sem PDMs após 3 tentativas: " + ", ".join(falhas))
    return resultados, arquivos


# =============================================================================
# LINHA DE COMANDO
# =============================================================================

def ler_codigos(caminho, tipo, por_classe):
    """(códigos, mapa classe -> códigos, coluna da classe) do arquivo de entrada.
    Aceita a coluna do tipo escolhido ou a genérica 'codigo'; o mapa só é
    montado com por_classe e quando o arquivo traz a coluna da classe."""
    df = (pd.read_excel(caminho) if caminho.lower().endswith(".xlsx")
          else pd.read_csv(caminho, sep=";"))
    col = next((c for c in (tipo, "codigo") if c in df.columns), None)
    if col is None:
        raise ValueError(f"o arquivo deve ter a coluna '{tipo}' (ou 'codigo'). "
                         f"Colunas do arquivo: {list(df.columns)}")
    codigos = df[col].dropna().astype(int).drop_duplicates().tolist()

    mapa, col_cl = {}, None
    if por_classe:
        col_cl = next((c for c in ("classe", "Classe", "codigoClasse")
                       if c in df.columns), None)
        if col_cl:
            val = df[[col, col_cl]].dropna()
            chave = (val[col_cl].astype(str).str.strip()
                     .str.replace(r"\.0$", "", regex=True))
            for cl, grupo in val.groupby(chave):
                mapa[cl] = grupo[col].astype(int).drop_duplicates().tolist()
    return codigos, mapa, col_cl


def _lista(texto, rotulo):
    """'1;2, 3' -> [1, 2, 3]"""
    partes = [p for p in re.split(r"[;,\s]+", texto.strip()) if p]
    invalidos = [p for p in partes if not p.isdigit()]
    if invalidos or not partes:
        raise ValueError(f"{rotulo} inválido(s): {', '.join(invalidos) or '(vazio)'}")
    return list(dict.fromkeys(int(p) for p in partes))


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(
        prog="extrator_comprasgov.py",
        description="Extrai os Registros de Preço do Compras.gov "
                    "(dadosabertos.compras.gov.br), sem interface gráfica.")
    ap.add_argument("arquivo", nargs="?",
                    help=".xlsx ou .csv (separador ;) com a coluna codigoItemCatalogo, "
                         "codigoPdm ou codigo")
    ap.add_argument("-c", "--codigos",
                    help='códigos avulsos separados por ";" (ex.: "451234;451235")')
    ap.add_argument("--classes",
                    help='classes separadas por ";": Classe -> PDMs -> CATMATs -> '
                         "registros, um arquivo e um relatório por classe")
    ap.add_argument("-t", "--tipo", type=str.upper, choices=("CATMAT", "PDM"),
                    default="CATMAT",
                    help="os códigos são CATMATs ou PDMs (parâmetro tipo da API). "
                         "Com --classes, PDM extrai direto por PDM, sem expandir os CATMATs")
    ap.add_argument("-i", "--inicio", default="", metavar="DD-MM-AAAA",
                    help="data da compra: início")
    ap.add_argument("-f", "--fim", default="", metavar="DD-MM-AAAA",
                    help="data da compra: fim")
    ap.add_argument("--formato", choices=("xlsx", "csv"), default="xlsx",
                    help="formato de saída (padrão: xlsx)")
    ap.add_argument("-s", "--saida", default=".", metavar="PASTA",
                    help="pasta de destino dos arquivos (padrão: a atual)")
    ap.add_argument("--por-classe", action="store_true",
                    help="um arquivo por classe: a classe vem da coluna 'classe' do "
                         "arquivo ou, sem ela, do codigoClasse dos registros")
    ap.add_argument("--corrompidas", metavar="PASTA",
                    help="guarda nesta pasta o CSV bruto das páginas com perda")
    ap.add_argument("--sem-descricao-bps", action="store_true",
                    help="mantém o descricaoItem do Compras.gov. Por padrão ele é "
                         "trocado pela descrição do BPS do mesmo CATMAT, como no "
                         "ExtratorCatmat com 'Usar a descrição do BPS' marcado")
    args = ap.parse_args(argv)
    descricao_bps = not args.sem_descricao_bps

    if sum(1 for x in (args.arquivo, args.codigos, args.classes) if x) != 1:
        ap.error("informe exatamente uma entrada: ARQUIVO, --codigos ou --classes")
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")   # console/arquivo fora do UTF-8

    tipo = TIPO_POR_ROTULO[args.tipo]
    d_ini, d_fim, err = validar_e_obter_datas(args.inicio, args.fim)
    if err:
        print(f"[ERRO] {err}")
        return 1

    mapa, col_cl = {}, None
    try:
        if args.classes:
            classes = [str(c) for c in _lista(args.classes, "Classe(s)")]
        elif args.codigos:
            codigos = _lista(args.codigos, "Código(s)")
        else:
            codigos, mapa, col_cl = ler_codigos(args.arquivo, tipo, args.por_classe)
            if not codigos:
                raise ValueError("nenhum código no arquivo.")
    except Exception as e:
        print(f"[ERRO] {e}")
        return 1

    pasta = args.saida
    os.makedirs(pasta, exist_ok=True)
    if args.corrompidas:
        os.makedirs(args.corrompidas, exist_ok=True)

    _log("Entrada : " + (f"classe(s) {' | '.join(classes)}" if args.classes
                          else f"{len(codigos)} código(s)"
                               + (f" de {args.arquivo}" if args.arquivo else "")))
    _log(f"Tipo    : {args.tipo} (tipo={tipo})")
    if d_ini or d_fim:
        fd = lambda s: "-".join(reversed(s.split("-"))) if s else "..."
        _log(f"Período : {fd(d_ini)} a {fd(d_fim)}")
    _log(f"Saída   : {os.path.abspath(pasta)} ({args.formato})")
    _log("Descrição: " + ("do BPS (descricaoItem trocado pelo texto do BPS; "
                          "CATMAT sem compra no BPS fica com o do Compras.gov)"
                          if descricao_bps else "do Compras.gov (--sem-descricao-bps)"))
    if args.por_classe and not args.classes:
        _log("Um arquivo por classe — origem da classe: "
             + (f"coluna '{col_cl}' do arquivo" if mapa else "campo codigoClasse dos registros"))
    _log("Ctrl+C cancela e salva o que já foi baixado.")
    _log()

    signal.signal(signal.SIGINT, _ao_ctrl_c)
    inicio = time.time()
    try:
        if args.classes:
            resultados, arquivos = fluxo_classes(classes, tipo, d_ini, d_fim, pasta,
                                                 args.formato, args.corrompidas,
                                                 descricao_bps)
        else:
            resultados, arquivos = fluxo_codigos(codigos, tipo, d_ini, d_fim, pasta,
                                                 args.formato, args.por_classe, mapa,
                                                 args.corrompidas, descricao_bps)
    except KeyboardInterrupt:
        _log("Interrompido.")
        return 130

    soma = lambda f: sum(f(r) for r in resultados)
    _log()
    _log("Extração cancelada — o que foi baixado está salvo." if _cancelado
         else "Extração concluída.")
    _log(f"Códigos processados   : {soma(lambda r: len(r.processados))} "
         f"de {soma(lambda r: len(r.codigos))}")
    _log(f"Registros gravados    : {_milhar(soma(lambda r: r.total))}")
    _log(f"Páginas reparadas     : {soma(lambda r: r.reparadas)}")
    _log(f"Páginas com perda     : {soma(lambda r: sum(len(p) for p in r.perdas.values()))}")
    _log(f"Códigos sem registros : {soma(lambda r: r.vazios)}")
    _log(f"Códigos com erro      : {soma(lambda r: len(r.erros))}")
    if descricao_bps:
        resumos = [b for r in resultados for b in r.bps.values()]
        sem_bps = sorted({c for b in resumos for c in b["sem_bps"]}, key=int)
        falhas  = sorted({c for b in resumos for c in b["falhas"]})
        _log(f"Descrição do BPS      : {_milhar(sum(b['trocadas'] for b in resumos))} "
             f"linha(s) trocada(s); {len(sem_bps)} CATMAT(s) sem compra no BPS "
             f"(ficou o texto do Compras.gov)")
        if falhas:
            _log(f"[ATENÇÃO] BPS sem resposta para {len(falhas)} CATMAT(s) — ficou o "
                 "texto do Compras.gov: " + ", ".join(falhas[:30])
                 + (" ..." if len(falhas) > 30 else ""))
    _log(f"Tempo                 : {time.time() - inicio:.0f} s")
    for a in arquivos:
        _log(f"  {a}")
    return 130 if _cancelado else 0


if __name__ == "__main__":
    sys.exit(main())
